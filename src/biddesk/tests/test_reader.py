"""Tests for biddesk.reader. Cached files only: nothing here touches the network."""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from biddesk import config, reader


# --------------------------------------------------------------- requirement extraction (pure)

def _page(text: str, page: int = 1) -> list[dict]:
    return [{"page": page, "text": text}]


def test_binding_verbs_are_caught_and_plain_sentences_are_not():
    pages = _page(
        "The Contractor shall provide all labor, supervision, and equipment required to clean "
        "the facility. This paragraph is purely informational and carries no obligation at all. "
        "Offerors must submit their quote by electronic mail to the Contracting Officer."
    )
    reqs = reader.extract_requirements(pages, "sow.pdf", "abcdefgh12345678")
    texts = [r.text for r in reqs]
    assert len(reqs) == 2
    assert texts[0].startswith("The Contractor shall provide all labor")
    assert texts[1].startswith("Offerors must submit")
    assert all("purely informational" not in t for t in texts)


def test_ids_carry_notice_source_and_page():
    text = "The Contractor shall maintain a quality control plan for the entire period."
    reqs = reader.extract_requirements(_page(text, page=7), "attachment_1.pdf", "1234567890abcdef")
    assert reqs[0].id == reader.requirement_id("1234567890abcdef", "attachment_1.pdf", 7, text)
    assert reqs[0].id.startswith("12345678-attachme-p7-")
    assert reqs[0].source_file == "attachment_1.pdf"
    assert reqs[0].page == 7


def test_an_id_is_a_hash_of_the_text_it_names_not_a_position():
    """D2: re-extraction cannot repoint a shipped citation at a different sentence."""
    first = "The Contractor shall maintain a quality control plan for the entire period."
    second = "The Contractor shall submit a monthly staffing report to the Government."
    both = reader.extract_requirements(
        _page(first + " " + second, page=7), "attachment_1.pdf", "1234567890abcdef")
    assert [r.text for r in both] == [first, second]        # document order is kept
    later = reader.extract_requirements(
        _page("An earlier sentence was inserted here by the amendment. " + first + " " + second,
              page=7),
        "attachment_1.pdf", "1234567890abcdef")
    by_text = {r.text: r.id for r in later}
    assert by_text[first] == both[0].id and by_text[second] == both[1].id


def test_the_same_sentence_in_two_files_gets_two_ids():
    text = "The Contractor shall maintain a quality control plan for the entire period."
    a = reader.extract_requirements(_page(text, page=1), "sow.pdf", "1234567890abcdef")[0]
    b = reader.extract_requirements(_page(text, page=1), "pws.pdf", "1234567890abcdef")[0]
    assert a.id != b.id


def test_line_wrapped_pdf_text_is_joined_before_splitting():
    wrapped = ("The Contractor shall furnish all personnel, equipment, tools, materials,\n"
               "supervision, and other items necessary to perform grounds maintenance\n"
               "services as defined in this Performance Work Statement.\n"
               "The Government will inspect the work weekly.")
    reqs = reader.extract_requirements(_page(wrapped), "pws.pdf", "n0000000")
    assert len(reqs) == 1
    assert "supervision, and other items necessary" in reqs[0].text
    assert "\n" not in reqs[0].text


def test_bullet_boundary_splits_two_requirements():
    bulleted = ("1. The Contractor shall remove all trash from occupied office space daily.\n"
                "2. The Contractor shall clean and disinfect all restroom fixtures each workday.\n")
    reqs = reader.extract_requirements(_page(bulleted), "sow.pdf", "n0000000")
    assert len(reqs) == 2
    assert "trash" in reqs[0].text and "restroom" in reqs[1].text


def test_short_sentences_are_dropped():
    reqs = reader.extract_requirements(_page("Offers must comply."), "x.pdf", "n0000000")
    assert reqs == []


def test_far_clause_listing_lines_are_dropped():
    listing = ("52.212-4 Contract Terms and Conditions; 52.219-6 Notice of Total Small Business "
               "Set-Aside; 52.222-41 Service Contract Labor Standards shall apply to this contract.")
    reqs = reader.extract_requirements(_page(listing), "clauses.pdf", "n0000000")
    assert reqs == []


def test_near_identical_text_is_deduped_case_and_whitespace_insensitive():
    dup = ("The Contractor shall submit a monthly quality control report to the COR.\n"
           "THE  CONTRACTOR   SHALL SUBMIT A MONTHLY QUALITY CONTROL REPORT TO THE COR.\n")
    reqs = reader.extract_requirements(_page(dup), "sow.pdf", "n0000000")
    assert len(reqs) == 1


def test_long_text_is_cut_at_a_clause_boundary_with_no_ellipsis():
    long_tail = "The Contractor shall provide " + ("cleaning services, floor care, " * 40) + "and reporting."
    reqs = reader.extract_requirements(_page(long_tail), "sow.pdf", "n0000000")
    t = reqs[0].text
    assert len(t) <= reader.MAX_REQ_CHARS
    assert not t.endswith(("...", "…"))
    assert t in reader._normalize(long_tail)  # still a verbatim prefix of the source sentence


@pytest.mark.parametrize("sentence,expected", [
    ("Offerors must submit their quote by email no later than the date shown on the form.", "submission"),
    ("Award shall be made on the basis of the lowest price technically acceptable proposal received.",
     "evaluation"),
    ("The period of performance shall be one base year plus four option years of service.", "schedule"),
    ("All personnel must possess a current state license and shall complete annual training.", "staffing"),
    ("The Contractor shall be registered in SAM and must comply with the wage determination.",
     "compliance"),
    ("The Contractor shall clean and disinfect all restroom fixtures each and every workday.", "scope"),
    ("The widget shall remain blue throughout the entirety of the referenced arrangement.", "other"),
])
def test_categories(sentence, expected):
    reqs = reader.extract_requirements(_page(sentence), "x.pdf", "n0000000")
    assert reqs and reqs[0].category == expected


# --------------------------------------------------------------- text extraction

def test_docx_keeps_paragraph_order_and_table_cells(tmp_path: Path):
    import docx
    d = docx.Document()
    d.add_paragraph("The Contractor shall provide janitorial services.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Frequency"
    t.cell(0, 1).text = "Daily"
    d.add_paragraph("End of statement.")
    p = tmp_path / "sow.docx"
    d.save(str(p))
    pages = reader.extract_text(p)
    assert len(pages) == 1 and pages[0]["page"] == 1
    text = pages[0]["text"]
    assert "Frequency | Daily" in text
    assert text.index("janitorial") < text.index("Frequency") < text.index("End of statement")


def test_xlsx_one_page_per_sheet_rows_joined_with_pipes(tmp_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "CLIN"
    ws.append(["Item", "Qty"])
    ws.append(["Mowing", 12])
    wb.create_sheet("Wages").append(["Janitor", "18.50"])
    p = tmp_path / "pricing.xlsx"
    wb.save(str(p))
    pages = reader.extract_text(p)
    assert [x["page"] for x in pages] == [1, 2]
    assert "Item | Qty" in pages[0]["text"] and "Mowing | 12" in pages[0]["text"]
    assert "Janitor | 18.5" in pages[1]["text"]


def test_unknown_binary_is_unsupported(tmp_path: Path):
    p = tmp_path / "legacy.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1binary")
    pages, scanned, unsupported = reader._extract(p)
    assert pages == [] and unsupported is True and scanned is False
    assert reader.extract_text(p) == []


def test_scanned_pdf_detection_flags_and_blanks_pages(monkeypatch, tmp_path: Path):
    p = tmp_path / "scan.pdf"
    p.write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(reader, "_pdf_pages",
                        lambda path: [{"page": 1, "text": "  "}, {"page": 2, "text": "3"}])
    pages, scanned, unsupported = reader._extract(p)
    assert scanned is True and unsupported is False
    assert pages == [{"page": 1, "text": ""}, {"page": 2, "text": ""}]


def test_text_pdf_is_not_flagged_scanned(monkeypatch, tmp_path: Path):
    p = tmp_path / "text.pdf"
    p.write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(reader, "_pdf_pages", lambda path: [{"page": 1, "text": "x" * 500}])
    pages, scanned, unsupported = reader._extract(p)
    assert scanned is False and unsupported is False and len(pages[0]["text"]) == 500


# --------------------------------------------------------------- cached-corpus checks

def _any_cached_pdf() -> Path | None:
    for d in sorted(config.ATTACH.iterdir()):
        if d.is_dir():
            for f in reader.cached_files(d.name):
                if f.suffix.lower() == ".pdf" and f.stat().st_size > 20_000:
                    return f
    return None


def test_real_cached_pdf_yields_numbered_pages():
    p = _any_cached_pdf()
    if p is None:
        pytest.skip("no cached PDF attachment; run `python -m biddesk.reader download` first")
    pages = reader.extract_text(p)
    assert pages and [x["page"] for x in pages] == list(range(1, len(pages) + 1))
    assert all(isinstance(x["text"], str) for x in pages)


def test_cached_files_excludes_sidecars_and_extracted_json():
    for d in sorted(config.ATTACH.iterdir()):
        if not d.is_dir():
            continue
        names = [f.name for f in reader.cached_files(d.name)]
        assert all(not n.endswith(".meta.json") for n in names)
        assert reader.EXTRACTED_NAME not in names
        return
    pytest.skip("no cached attachments")


def test_resolve_document_skips_the_meta_sidecar(tmp_path: Path):
    (tmp_path / "abc.meta.json").write_text("{}", encoding="utf-8")
    (tmp_path / "abc.pdf").write_bytes(b"%PDF-1.4")
    assert reader._resolve_document(tmp_path / "abc.meta.json").name == "abc.pdf"
    assert reader._resolve_document(tmp_path / "abc.pdf").name == "abc.pdf"


def test_extracted_json_shape_if_present():
    found = None
    for d in sorted(config.ATTACH.iterdir()):
        if d.is_dir() and (d / reader.EXTRACTED_NAME).exists():
            found = json.loads((d / reader.EXTRACTED_NAME).read_text(encoding="utf-8"))
            break
    if found is None:
        pytest.skip("no extracted.json yet; run `python -m biddesk.reader extract` first")
    assert set(found) >= {"notice_id", "files", "requirements", "extracted_at"}
    for f in found["files"]:
        assert set(f) >= {"path", "original_name", "pages", "chars", "scanned", "unsupported"}
    for r in found["requirements"]:
        assert set(r) >= {"id", "text", "source_file", "page", "category"}
        assert len(r["text"]) <= reader.MAX_REQ_CHARS
