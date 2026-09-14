"""Attachment reader: download every solicitation attachment, extract per-page text,
and lift binding requirement sentences deterministically. No model calls anywhere in this file.

Usage:
    python -m biddesk.reader download [--profiles red-cedar ...] [--max-per-notice 8]
    python -m biddesk.reader extract  [--profiles red-cedar ...]
    python -m biddesk.reader report

Writes data/attachments/_download_report.json and data/attachments/<notice_id>/extracted.json.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from zoneinfo import ZoneInfo

from . import config
from .models import Notice, Requirement
from .profiles import PROFILES, REAL_SOLICITATION_TYPES
from .sam_client import SamError, fetch_attachment
from .snapshot import load as load_snapshot

CENTRAL = ZoneInfo("America/Chicago")

DOWNLOAD_REPORT = config.ATTACH / "_download_report.json"
EXTRACTED_NAME = "extracted.json"

# Suffixes we can turn into text. Everything else is recorded as unsupported.
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLS_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
TEXT_SUFFIXES = {".txt", ".md", ".csv"}
# URL-level skip list. Every SAM resource link is an extension-less /download endpoint,
# so this only fires if a direct media/archive link ever shows up in a snapshot.
SKIP_URL_SUFFIXES = {".zip", ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff",
                     ".mp4", ".mov", ".avi", ".exe", ".msi", ".dwg"}

SCANNED_CHAR_THRESHOLD = 200
MAX_REQ_CHARS = 600
MIN_REQ_CHARS = 22


def now_iso() -> str:
    return dt.datetime.now(CENTRAL).isoformat(timespec="seconds")


# --------------------------------------------------------------------------- download

def _real_notices(slugs: list[str] | None) -> list[Notice]:
    """Every notice of a real solicitation type across the named snapshots, deduped by notice_id."""
    slugs = list(slugs) if slugs else list(PROFILES)
    seen: dict[str, Notice] = {}
    for slug in slugs:
        _meta, notices = load_snapshot(slug)
        for n in notices:
            if n.type in REAL_SOLICITATION_TYPES and n.notice_id not in seen:
                seen[n.notice_id] = n
    return list(seen.values())


def _is_document_url(url: str) -> bool:
    return Path(url.split("?")[0].rstrip("/")).suffix.lower() not in SKIP_URL_SUFFIXES


def _resolve_document(path: Path) -> Path:
    """sam_client.fetch_attachment globs `<fid>.*` on a cache hit, which also matches the
    `<fid>.meta.json` sidecar it writes. Re-resolve to the real payload when that happens."""
    if path.name.endswith(".meta.json"):
        stem = path.name[: -len(".meta.json")]
        for cand in sorted(path.parent.glob(f"{stem}.*")):
            if not cand.name.endswith(".meta.json"):
                return cand
    return path


def download_all(slugs=None, max_per_notice: int = 8) -> dict:
    """Download every resource link of every real-solicitation notice. Errors are recorded, not raised."""
    notices = _real_notices(slugs)
    jobs: list[tuple[str, str]] = []
    for n in notices:
        for url in n.resource_links[:max_per_notice]:
            if _is_document_url(url):
                jobs.append((n.notice_id, url))
    print(f"{len(notices)} real-type notices, {len(jobs)} links to fetch "
          f"(cap {max_per_notice} per notice)", file=sys.stderr)

    def one(job: tuple[str, str]):
        notice_id, url = job
        try:
            return notice_id, str(_resolve_document(fetch_attachment(notice_id, url)))
        except SamError as e:
            return notice_id, f"ERROR: {e}"
        except Exception as e:  # requests timeouts, connection resets: record and keep going
            return notice_id, f"ERROR: {type(e).__name__}: {e}"

    summary: dict[str, list[str]] = {n.notice_id: [] for n in notices}
    with ThreadPoolExecutor(8) as ex:
        for notice_id, result in ex.map(one, jobs):
            summary[notice_id].append(result)

    ok = sum(1 for v in summary.values() for r in v if not r.startswith("ERROR:"))
    bad = sum(1 for v in summary.values() for r in v if r.startswith("ERROR:"))
    DOWNLOAD_REPORT.write_text(json.dumps(
        {"written_at": now_iso(), "notices": len(summary), "links": len(jobs),
         "downloaded": ok, "failed": bad, "max_per_notice": max_per_notice,
         "results": summary}, indent=1), encoding="utf-8")
    print(f"downloaded {ok}, failed {bad} -> {DOWNLOAD_REPORT}", file=sys.stderr)
    return summary


# --------------------------------------------------------------------------- text extraction

def _pdf_pages(path: Path) -> list[dict]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        pages.append({"page": i, "text": txt})
    return pages


def _docx_pages(path: Path) -> list[dict]:
    """Whole document is page 1. Body elements are walked in document order so paragraphs and
    tables keep their relative position; table cells are joined with ' | '."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    chunks: list[str] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                chunks.append(t)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    chunks.append(line)
    return [{"page": 1, "text": "\n".join(chunks)}]


def _xlsx_pages(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            # A worksheet row is a hard sentence boundary: spreadsheet cells carry no
            # terminal punctuation, so without this every row merges into one segment.
            pages.append({"page": i, "text": "\n\x00".join(rows)})
    finally:
        wb.close()
    return pages


def _text_pages(path: Path) -> list[dict]:
    return [{"page": 1, "text": path.read_text(encoding="utf-8", errors="replace")}]


def _extract(path: Path) -> tuple[list[dict], bool, bool]:
    """Returns (pages, scanned, unsupported)."""
    suffix = path.suffix.lower()
    try:
        if suffix in PDF_SUFFIXES:
            pages = _pdf_pages(path)
            total = sum(len(p["text"].strip()) for p in pages)
            if pages and total < SCANNED_CHAR_THRESHOLD:
                # Image-only PDF. The multimodal path reads these later; no OCR here.
                return [{"page": p["page"], "text": ""} for p in pages], True, False
            return pages, False, False
        if suffix in DOCX_SUFFIXES:
            return _docx_pages(path), False, False
        if suffix in XLS_SUFFIXES:
            return _xlsx_pages(path), False, False
        if suffix in TEXT_SUFFIXES:
            return _text_pages(path), False, False
    except Exception as e:  # encrypted PDF, legacy .xls, corrupt payload
        print(f"  [extract] {path.name}: {type(e).__name__}: {e}", file=sys.stderr)
        return [], False, True
    return [], False, True


def extract_text(path) -> list[dict]:
    """Per-page text: [{"page": 1-based int, "text": str}, ...]. Empty list when unsupported."""
    return _extract(Path(path))[0]


# --------------------------------------------------------------------------- requirements

BINDING = re.compile(
    r"\bshall\b|\bmust\b|\bis required to\b|\bare required to\b|\bwill be required\b"
    r"|\bcontractor will\b|\boffers must\b|\bquotes shall\b", re.I)

CLAUSE_NUM = re.compile(r"\b\d{2}\.\d{3}(?:-\d{1,2})?(?:\s*Alt\s*[IVX]+)?\b")
# U+F0B7 U+F0A7 U+F0D8 U+F06C U+F0FC are the Wingdings/Symbol private-use glyphs Word writes for
# bullets; pypdf hands them through, and without them a bulleted task list never splits.
BULLET_LINE = re.compile(r"^\s*(?:[•‣▪●·–—*o]\s+"
                         r"|\(?[a-zA-Z]\)\s+|\(?\d{1,2}(?:\.\d{1,2}){0,3}[.)]\s+"
                         r"|\d{1,2}(?:\.\d{1,2}){1,3}\s+)")   # "2.1 Factor1", "5.1.3 Eligibility"
SPLIT = re.compile(r"(?<=\.)\s+(?=[A-Z])|\x00")
CUT_POINTS = (";", ",", ":", " and ", " or ")

CATEGORY_RULES: list[tuple[str, list[str]]] = [
    ("submission", [r"\bquot", r"\bproposal", r"\bsubmit", r"\bdue\b", r"\bemail\b",
                    r"offeror\s+shall\s+provide"]),
    ("evaluation", [r"evaluat", r"\bawarde?d?\b", r"\bbasis\b", r"\bfactors?\b",
                    r"\bLPTA\b", r"best value"]),
    ("schedule", [r"period of performance", r"\bhours?\b", r"\bdays?\b", r"\bdeadline",
                  r"within\s+\w+\s+(?:calendar |business |working |government )?days?",
                  r"base year", r"option year"]),
    ("staffing", [r"personnel", r"employee", r"\bstaff", r"key personnel", r"certif",
                  r"licens", r"background", r"clearance", r"training"]),
    ("compliance", [r"\bFAR\b", r"\bDFARS\b", r"wage determination", r"insurance",
                    r"SAM\s+registration", r"registered in SAM", r"e-?Verify", r"\bOSHA\b",
                    r"\bHIPAA\b", r"security"]),
    ("scope", [r"perform", r"provide", r"service", r"\bclean", r"maintain", r"support",
               r"deliver"]),
]
_COMPILED_RULES = [(name, [re.compile(p, re.I) for p in pats]) for name, pats in CATEGORY_RULES]


def categorize(text: str) -> str:
    """Most keyword hits wins; ties go to the earlier rule. A single incidental word
    ('proposal' inside an award-basis sentence) must not outrank the dominant category."""
    best, best_score = "other", 0
    for name, pats in _COMPILED_RULES:
        score = sum(1 for p in pats if p.search(text))
        if score > best_score:
            best, best_score = name, score
    return best


def _is_clause_listing(sentence: str) -> bool:
    """A FAR/DFARS clause-listing line: mostly clause numbers like 52.212-4, 252.204-7012."""
    nums = CLAUSE_NUM.findall(sentence)
    if not nums:
        return False
    if CLAUSE_NUM.match(sentence.lstrip()):   # the line opens with its clause number
        return True
    if len(nums) >= 3:
        return True
    covered = sum(len(n) for n in nums)
    return covered / max(len(sentence), 1) > 0.25


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _trim(text: str) -> str:
    """Verbatim, whitespace-normalized, at most MAX_REQ_CHARS, cut at a clause boundary, no ellipsis."""
    text = _normalize(text)
    if len(text) <= MAX_REQ_CHARS:
        return text
    window = text[:MAX_REQ_CHARS]
    best = max((window.rfind(c) for c in CUT_POINTS), default=-1)
    if best < MAX_REQ_CHARS // 3:
        best = window.rfind(" ")
    if best <= 0:
        return window.rstrip()
    return window[:best].rstrip().rstrip(",;:")


def _sentences(page_text: str) -> list[str]:
    """Join PDF line-wrapping, then split on sentence, semicolon, and bullet/numbered boundaries."""
    if not page_text:
        return []
    text = page_text.replace("\r\n", "\n").replace("\r", "\n")
    # De-hyphenate a word broken across a line, but keep a real hyphen when the continuation
    # is not a lowercase letter ("a 1-to-\n2 sentence" must not become "a 1-to2 sentence").
    text = re.sub(r"(\w)-\n(\w)",
                  lambda m: m.group(1) + m.group(2) if m.group(2).islower()
                  else m.group(1) + "-" + m.group(2),
                  text)
    lines = text.split("\n")
    joined_parts = []
    for line in lines:
        if BULLET_LINE.match(line) or (joined_parts and not line.strip()):
            joined_parts.append("\x00")
        joined_parts.append(line.strip())
    joined = " ".join(joined_parts)
    joined = re.sub(r"[ \t]+", " ", joined)
    out = []
    for seg in SPLIT.split(joined):
        seg = _normalize(seg)
        if seg:
            out.append(seg)
    return out


def _dedupe_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def requirement_id(notice_id: str, source_file: str, page: int, text: str) -> str:
    """Content-addressed requirement id.

    ``<notice8>-<file8>-p<page>-<sha1(source_file + "|" + normalized text)[:8]>``. The id is a
    function of the text it names, so re-extracting an attachment cannot repoint a shipped
    citation at a different sentence: either the same sentence is found again under the same
    id, or the id is gone and :func:`desk.validate_matrix` drops the row.
    """
    stem = Path(source_file).stem or source_file
    digest = hashlib.sha1((source_file + "|" + _normalize(text)).encode("utf-8")).hexdigest()[:8]
    return f"{notice_id[:8]}-{stem[:8]}-p{int(page)}-{digest}"


def extract_requirements(pages: list[dict], source_file: str, notice_id: str = "") -> list[Requirement]:
    """Deterministic, regex-only requirement lift, in document order. No model call."""
    out: list[Requirement] = []
    seen: set[str] = set()
    for page in pages:
        pno = int(page.get("page", 0))
        for sentence in _sentences(page.get("text") or ""):
            if not BINDING.search(sentence):
                continue
            if _is_clause_listing(sentence):
                continue
            text = _trim(sentence)
            if len(text) < MIN_REQ_CHARS:
                continue
            key = _dedupe_key(text)
            if key in seen:
                continue
            seen.add(key)
            out.append(Requirement(
                id=requirement_id(notice_id, source_file, pno, text),
                text=text,
                source_file=source_file,
                page=pno,
                category=categorize(text),
            ))
    return out


# --------------------------------------------------------------------------- per-notice read

def cached_files(notice_id: str) -> list[Path]:
    """Real attachment payloads for a notice: no sidecars, no extracted.json, no directories."""
    d = config.ATTACH / notice_id
    if not d.is_dir():
        return []
    return sorted(p for p in d.iterdir()
                  if p.is_file() and not p.name.endswith(".meta.json") and p.name != EXTRACTED_NAME)


def _original_name(path: Path) -> str:
    meta = path.with_suffix("")
    meta = meta.parent / f"{meta.name}.meta.json"
    if meta.exists():
        try:
            return json.loads(meta.read_text(encoding="utf-8")).get("original_name") or path.name
        except Exception:
            pass
    return path.name


def read_notice(notice: Notice) -> dict:
    """Extract text and requirements for every cached attachment plus the notice description."""
    files: list[dict] = []
    reqs: list[Requirement] = []
    for path in cached_files(notice.notice_id):
        pages, scanned, unsupported = _extract(path)
        chars = sum(len(p["text"]) for p in pages)
        files.append({"path": str(path), "original_name": _original_name(path),
                      "pages": len(pages), "chars": chars,
                      "scanned": scanned, "unsupported": unsupported})
        if pages and not scanned and not unsupported:
            reqs.extend(extract_requirements(pages, path.name, notice.notice_id))
    if notice.description_text:
        reqs.extend(extract_requirements(
            [{"page": 0, "text": notice.description_text}], "description", notice.notice_id))
    ids = [r.id for r in reqs]
    if len(ids) != len(set(ids)):
        dupes = sorted({i for i in ids if ids.count(i) > 1})
        raise ValueError(f"duplicate requirement ids for {notice.notice_id}: {dupes[:5]}")
    out = {"notice_id": notice.notice_id,
           "title": notice.title,
           "files": files,
           "requirements": [r.model_dump() for r in reqs],
           "extracted_at": now_iso()}
    d = config.ATTACH / notice.notice_id
    d.mkdir(parents=True, exist_ok=True)
    (d / EXTRACTED_NAME).write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
    return out


def extract_all(slugs=None, resume: bool = False) -> list[dict]:
    """resume=True skips notices that already have an extracted.json, so a long run that was
    interrupted does not re-parse the large scanned PDFs it already got through."""
    outs = []
    for n in _real_notices(slugs):
        if not cached_files(n.notice_id) and not n.description_text:
            continue
        if resume and (config.ATTACH / n.notice_id / EXTRACTED_NAME).exists():
            continue
        outs.append(read_notice(n))
    return outs


# --------------------------------------------------------------------------- report

def _ascii(s: str, width: int) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()[:width]
    return s.encode("ascii", "replace").decode("ascii")


def report() -> dict:
    rows = []
    for d in sorted(config.ATTACH.iterdir()):
        if not d.is_dir():
            continue
        f = d / EXTRACTED_NAME
        if not f.exists():
            continue
        e = json.loads(f.read_text(encoding="utf-8"))
        rows.append({
            "notice_id": e["notice_id"],
            "title": e.get("title", ""),
            "files": len(e["files"]),
            "pages": sum(x["pages"] for x in e["files"]),
            "requirements": len(e["requirements"]),
            "scanned": sum(1 for x in e["files"] if x["scanned"]),
            "unsupported": sum(1 for x in e["files"] if x["unsupported"]),
        })
    rows.sort(key=lambda r: (-r["requirements"], r["notice_id"]))
    hdr = f"{'notice_id':34s} {'title':50s} {'files':>5s} {'pages':>6s} {'reqs':>5s} {'scan':>4s} {'uns':>4s}"
    print(hdr)
    print("-" * len(hdr))
    for r in rows:
        print(f"{r['notice_id']:34s} {_ascii(r['title'], 50):50s} {r['files']:5d} "
              f"{r['pages']:6d} {r['requirements']:5d} {r['scanned']:4d} {r['unsupported']:4d}")
    totals = {
        "notices": len(rows),
        "notices_with_attachment": sum(1 for r in rows if r["files"]),
        "files": sum(r["files"] for r in rows),
        "pages": sum(r["pages"] for r in rows),
        "requirements": sum(r["requirements"] for r in rows),
        "scanned_files": sum(r["scanned"] for r in rows),
        "unsupported_files": sum(r["unsupported"] for r in rows),
    }
    print("-" * len(hdr))
    print("TOTALS  " + "  ".join(f"{k}={v}" for k, v in totals.items()))
    return {"rows": rows, "totals": totals}


# --------------------------------------------------------------------------- CLI

def main(argv=None):
    ap = argparse.ArgumentParser(prog="python -m biddesk.reader")
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("download")
    d.add_argument("--profiles", nargs="*", default=None)
    d.add_argument("--max-per-notice", type=int, default=8)
    x = sub.add_parser("extract")
    x.add_argument("--profiles", nargs="*", default=None)
    x.add_argument("--resume", action="store_true", help="skip notices that already have extracted.json")
    sub.add_parser("report")
    a = ap.parse_args(argv)
    if a.cmd == "download":
        download_all(a.profiles, a.max_per_notice)
    elif a.cmd == "extract":
        outs = extract_all(a.profiles, resume=a.resume)
        print(f"extracted {len(outs)} notices, "
              f"{sum(len(o['requirements']) for o in outs)} requirements", file=sys.stderr)
    else:
        report()


if __name__ == "__main__":
    main()
